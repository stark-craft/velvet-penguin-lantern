"""Exports of the current edited snapshot; no AI calls on export."""
import base64
from io import BytesIO
from html import escape
from pathlib import Path
import os
from bs4 import BeautifulSoup
from fastapi import HTTPException
from core.settings import PROJECT_ROOT
from .document import sanitize_html, export_html, safe_url


def image_data(url):
    if url.startswith('data:image/'):
        try:
            raw = base64.b64decode(url.split(',', 1)[1], validate=True)
            from PIL import Image
            image = Image.open(BytesIO(raw))
            if image.width * image.height > 40000000:
                return None
            out = BytesIO()
            image.convert('RGB').save(out, format='PNG')
            out.seek(0)
            return out
        except Exception:
            return None
    # Reuse the established public-URL, redirect, TLS and size protections.
    from news_scrapper.application import download_image_for_export
    return download_image_for_export(url) if url else None


def blocks(soup):
    for node in soup.find_all(['h1', 'h2', 'h3', 'h4', 'p', 'li', 'blockquote', 'img', 'hr', 'div']):
        if node.name == 'div' and node.find(['h1','h2','h3','h4','p','li','blockquote','img','hr','div']):
            continue
        if node.name in {'p', 'div'} and node.find_parent(['li','blockquote']):
            continue
        yield node


def text_of(node):
    return node.get_text(' ', strip=True)


def pdf_inline(node):
    if isinstance(node, str):
        return escape(node)
    contents = ''.join(pdf_inline(child) for child in node.children)
    name = {'strong':'b','em':'i'}.get(node.name, node.name)
    if name in {'b','i','u','s'}:
        return f'<{name}>{contents}</{name}>'
    if name == 'br':
        return '<br/>'
    if name == 'a' and node.get('href'):
        return f'<a href="{escape(node["href"], quote=True)}">{contents}</a>'
    return contents


def pdf_fonts(styles):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    paths = [Path(os.getenv('WINDIR', 'C:/Windows'))/'Fonts'/'malgun.ttf', Path('/System/Library/Fonts/Supplemental/Arial Unicode.ttf'), Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')]
    font = next((p for p in paths if p.exists()), None)
    if font:
        if 'SamparkReport' not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont('SamparkReport', str(font)))
            pdfmetrics.registerFontFamily('SamparkReport', normal='SamparkReport', bold='SamparkReport', italic='SamparkReport', boldItalic='SamparkReport')
        for style in styles.byName.values():
            style.fontName = 'SamparkReport'


def export_snapshot(fmt, title, html):
    soup = BeautifulSoup(sanitize_html(html), 'html.parser')
    out = BytesIO()
    if fmt == 'html':
        return export_html(str(soup), title).encode(), 'text/html', 'html'
    if fmt == 'pdf':
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, PageBreak
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.pagesizes import A4, landscape
        styles = getSampleStyleSheet()
        pdf_fonts(styles)
        flow = []
        for node in blocks(soup):
            if node.name == 'hr' and node.get('data-kind') == 'page-break':
                flow.append(PageBreak())
            elif node.name == 'img':
                data = image_data(node.get('src', ''))
                if data:
                    img = Image(data)
                    factor = min(280 / img.imageWidth, 190 / img.imageHeight)
                    img.drawWidth, img.drawHeight = img.imageWidth * factor, img.imageHeight * factor
                    flow.extend([img, Spacer(1, 10)])
            else:
                txt = pdf_inline(node) or '<br/>'
                if node.name == 'li':
                    txt = '• ' + txt
                style = styles[{'h1':'Title', 'h2':'Heading1', 'h3':'Heading2', 'h4':'Heading3'}.get(node.name, 'BodyText')]
                flow.extend([Paragraph(txt, style), Spacer(1, 8)])
        SimpleDocTemplate(out, pagesize=landscape(A4), title=title, leftMargin=40, rightMargin=40, topMargin=36, bottomMargin=36).build(flow)
        return out.getvalue(), 'application/pdf', 'pdf'
    if fmt == 'docx':
        from docx import Document
        from docx.shared import Inches
        doc = Document()
        section = doc.sections[0]
        section.page_width, section.page_height = Inches(11.7), Inches(8.3)
        for node in blocks(soup):
            if node.name == 'hr' and node.get('data-kind') == 'page-break':
                doc.add_page_break()
            elif node.name == 'img':
                data = image_data(node.get('src', ''))
                if data:
                    doc.add_picture(data, width=Inches(3.5))
            elif node.name.startswith('h'):
                doc.add_heading(text_of(node), level=min(int(node.name[1]), 4))
            else:
                para = doc.add_paragraph(style='List Bullet' if node.name == 'li' else None)
                for child in node.descendants:
                    if isinstance(child, str):
                        run = para.add_run(child)
                        parents = [p.name for p in child.parents if p is not node]
                        run.bold = bool({'b','strong'} & set(parents))
                        run.italic = bool({'i','em'} & set(parents))
                        run.underline = 'u' in parents
        doc.save(out)
        return out.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'docx'
    if fmt == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        book = Workbook()
        sheet = book.active
        sheet.title = 'Edited report'
        sheet.append(['Section', 'Content', 'Source URL'])
        for node in blocks(soup):
            if node.name in {'img', 'hr'}:
                continue
            parent = node.find_parent('section')
            link = node.find('a')
            values = [(parent or {}).get('data-kind', 'report'), text_of(node), link.get('href','') if link else '']
            # Untrusted article text must not become spreadsheet formulas.
            sheet.append(["'"+v if v.startswith(('=', '+', '-', '@')) else v for v in values])
        sheet.column_dimensions['A'].width = 18
        sheet.column_dimensions['B'].width = 100
        sheet.column_dimensions['C'].width = 55
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        book.save(out)
        return out.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'xlsx'
    if fmt == 'pptx':
        return export_ppt(title, soup)
    raise HTTPException(422, 'Unsupported export format.')



def _pptx_article(section):
    """Read the edited article snapshot; never regenerate or guess its content."""
    heading = section.find(['h1', 'h2'])
    meta = heading.find_next_sibling('p') if heading else None
    metadata = text_of(meta) if meta else ''
    source, sep, date = metadata.rpartition(' · ')
    if not sep:
        source, date = metadata, ''

    def from_area(kind):
        area = section.find(attrs={'data-kind': kind})
        return ' '.join(text_of(node) for node in area.find_all(['p', 'li']) if not node.find_parent('li')) if area else ''

    def after_heading(label):
        h3 = next((h for h in section.find_all('h3') if label in text_of(h).casefold()), None)
        if h3 is None:
            return ''
        parts = []
        for node in h3.next_siblings:
            if getattr(node, 'name', None) in ('h1', 'h2', 'h3', 'h4'):
                break
            if getattr(node, 'name', None) in ('p', 'li') and not node.find('a', href=True):
                parts.append(text_of(node))
            elif getattr(node, 'name', None) == 'div':
                parts.extend(text_of(p) for p in node.find_all(['p', 'li']) if not p.find('a', href=True) and not p.find_parent('li'))
        return ' '.join(parts)

    links = [a for a in section.find_all('a', href=True) if safe_url(a.get('href'))]
    preferred = next((a for a in links if 'original article' in text_of(a).casefold()), None)
    link = safe_url((preferred or (links[0] if links else None)).get('href')) if links else ''
    return {
        'title': text_of(heading) if heading else 'Untitled article',
        'source': source.strip(),
        'date': date.strip(),
        'summary': from_area('article-summary') or after_heading('summary'),
        'insight': from_area('impact-body') or after_heading('samsung') or after_heading('matter'),
        'team': after_heading('team'),
        'link': link,
    }


def _pptx_template_run(source):
    if source is not None and source.has_text_frame:
        for para in source.text_frame.paragraphs:
            if para.runs:
                return para.runs[0]
    return None


def _pptx_font_pt(source, fallback):
    run = _pptx_template_run(source)
    if run is not None and run.font.size is not None:
        return run.font.size.pt
    return fallback


def _pptx_text_width(text, font_pt, bold=False):
    """Conservative point-width estimate; prefer installed Calibri via Pillow."""
    try:
        from PIL import ImageFont
        import os
        fonts = ['calibrib.ttf' if bold else 'calibri.ttf', 'arialbd.ttf' if bold else 'arial.ttf', 'DejaVuSans-Bold.ttf' if bold else 'DejaVuSans.ttf']
        for name in fonts:
            for candidate in (os.path.join(os.environ.get('WINDIR', 'C:/Windows'), 'Fonts', name), name):
                try:
                    font = ImageFont.truetype(candidate, max(1, round(font_pt * 4)))
                    return font.getlength(text) / 4
                except OSError:
                    continue
    except ImportError:
        pass
    return sum(1.0 if ord(ch) > 0x2e80 else 0.65 if ch.isupper() else 0.54 for ch in text) * font_pt


def _pptx_fits(text, shape, pt, bold=False):
    import re
    frame = shape.text_frame
    # EMU to points: 12,700 EMU per point.
    width = max(0, (shape.width - frame.margin_left - frame.margin_right) / 12700 * 0.96)
    height = max(0, (shape.height - frame.margin_top - frame.margin_bottom) / 12700 * 0.92)
    max_lines = max(1, int(height / (pt * 1.18)))
    if not width:
        return False
    lines, line = 1, ''
    for word in re.findall(r'\S+', text):
        proposed = (line + ' ' + word).strip()
        if _pptx_text_width(proposed, pt, bold) <= width:
            line = proposed
        elif line and _pptx_text_width(word, pt, bold) <= width:
            lines, line = lines + 1, word
        else:
            return False  # Overwide unbreakable token; truncate, don't overflow.
        if lines > max_lines:
            return False
    return True


def _pptx_fit_text(text, shape, pt, bold=False, suffix=''):
    """Keep the largest word-boundary prefix fitting the fixed placeholder."""
    import re
    clean = re.sub(r'\s+', ' ', str(text or '')).strip()
    if _pptx_fits(clean + suffix, shape, pt, bold):
        return clean
    words = clean.split(' ')
    low, high, best = 0, len(words), ''
    while low <= high:
        mid = (low + high) // 2
        candidate = ' '.join(words[:mid]).rstrip(' ,;:') + '…'
        if _pptx_fits(candidate + suffix, shape, pt, bold):
            best, low = candidate, mid + 1
        else:
            high = mid - 1
    return best or ('…' if _pptx_fits('…' + suffix, shape, pt, bold) else '')


def _pptx_set_placeholder(shape, source, text, hyperlink=''):
    """Copy the marker run styling instead of destroying template formatting."""
    from copy import deepcopy
    from pptx.enum.text import MSO_AUTO_SIZE
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    para = frame.paragraphs[0]
    template_run = _pptx_template_run(source)
    if source is not None and source.has_text_frame:
        original_para = source.text_frame.paragraphs[0]
        if original_para._p.pPr is not None:
            para._p.insert(0, deepcopy(original_para._p.pPr))
    run = para.add_run()
    run.text = text
    if template_run is not None and template_run._r.rPr is not None:
        run._r.insert(0, deepcopy(template_run._r.rPr))
    if hyperlink:
        run.hyperlink.address = hyperlink
    return run


def _pptx_set_headline(shape, source, headline, publisher, link):
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    import re
    headline = re.sub(r'\s+', ' ', headline).strip()
    publisher = re.sub(r'\s+', ' ', publisher).strip()[:80]
    byline = f' By {publisher}' if publisher else ''
    size = next((pt for pt in (24, 22, 20, 18) if _pptx_fits(headline + byline, shape, pt, True)), 18)
    headline = _pptx_fit_text(headline, shape, size, True, byline)
    run = _pptx_set_placeholder(shape, source, headline, link)
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 112, 192)
    if byline:
        following = shape.text_frame.paragraphs[0].add_run()
        following.text = byline
        following.font.name = 'Calibri'
        following.font.size = Pt(size)
        following.font.bold = False
        following.font.color.rgb = RGBColor(0, 0, 0)


def export_ppt(title, soup):
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER
    path = PROJECT_ROOT / 'template.pptx'
    if not path.is_file():
        raise HTTPException(409, 'Add template.pptx to the application root to export PowerPoint.')
    deck = Presentation(str(path))
    layouts = {layout.name: layout for layout in deck.slide_layouts}
    if 'NewsLayoutFirst' not in layouts or 'NewsLayout' not in layouts:
        raise HTTPException(409, 'template.pptx needs both NewsLayoutFirst and NewsLayout layouts.')
    articles = soup.select('section[data-kind="article"]')
    if not articles:
        raise HTTPException(422, 'Select at least one article to export as PowerPoint.')
    # Sample slides are not user report content.
    while deck.slides:
        slide_id = deck.slides._sldIdLst[0]
        deck.part.drop_rel(slide_id.rId)
        deck.slides._sldIdLst.remove(slide_id)
    for index, section in enumerate(articles):
        layout = layouts['NewsLayoutFirst' if index == 0 else 'NewsLayout']
        slide = deck.slides.add_slide(layout)
        article = _pptx_article(section)
        image_node = section.find('img')
        for shape in slide.placeholders:
            if shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE:
                data = image_data(image_node.get('src', '')) if image_node else None
                if data and hasattr(shape, 'insert_picture'):
                    shape.insert_picture(data)
                continue
            if not shape.has_text_frame:
                continue
            source = layout.placeholders.get(shape.placeholder_format.idx)
            marker = (source.text if source is not None and source.has_text_frame else shape.text).strip()
            if '#TITLE' in marker:
                _pptx_set_headline(shape, source, article['title'], article['source'], article['link'])
                continue
            mapping = {
                '#SUMMARY': ('summary', 16),
                '#INSIGHT': ('insight', 10.5),
                '#Targated_SRID_Team': ('team', 12),
                '#DATE_HERE': ('date', 10),
                '#LINK': ('link', 10),
            }
            for token, (field, fallback_pt) in mapping.items():
                if token not in marker:
                    continue
                value = article[field]
                if token == '#LINK':
                    value = 'Open original article' if value else ''
                pt = _pptx_font_pt(source, fallback_pt)
                value = _pptx_fit_text(value, shape, pt)
                _pptx_set_placeholder(shape, source, value, article['link'] if token == '#LINK' else '')
                break
    out = BytesIO()
    deck.save(out)
    return out.getvalue(), 'application/vnd.openxmlformats-officedocument.presentationml.presentation', 'pptx'
