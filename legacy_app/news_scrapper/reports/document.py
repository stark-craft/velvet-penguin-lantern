"""Report HTML is the edited source of truth, never regenerated during export."""
import re
from urllib.parse import urlsplit
from bs4 import BeautifulSoup
from core.settings import PROJECT_ROOT

TAGS = {'section', 'div', 'p', 'br', 'h1', 'h2', 'h3', 'h4', 'strong', 'b', 'em', 'i', 'u', 's', 'ul', 'ol', 'li', 'blockquote', 'a', 'img', 'span', 'hr', 'font'}


def safe_url(value, image=False):
    value = str(value or '').strip()
    if image and re.fullmatch(r'data:image/(?:png|jpeg|webp);base64,[A-Za-z0-9+/=]+', value):
        return value
    try:
        parsed = urlsplit(value)
        return value if parsed.scheme in {'http', 'https'} and parsed.hostname and not parsed.username else ''
    except ValueError:
        return ''


def sanitize_html(html):
    soup = BeautifulSoup(html, 'html.parser')
    for tag in list(soup.find_all(True)):
        if tag.name is None:
            continue
        if tag.name in {'script', 'style', 'iframe', 'object', 'svg', 'form', 'input'}:
            tag.decompose()
        elif tag.name not in TAGS:
            tag.unwrap()
        else:
            attrs = {}
            if tag.get('data-kind') in {'overview', 'analysis', 'article', 'page-break', 'samsung-impact', 'impact-body', 'cross-analysis', 'article-summary'}:
                attrs['data-kind'] = tag['data-kind']
            if tag.name == 'a':
                attrs.update(href=safe_url(tag.get('href')), rel='noopener noreferrer', target='_blank')
            if tag.name == 'img':
                attrs.update(src=safe_url(tag.get('src'), image=True), alt=str(tag.get('alt', ''))[:300], referrerpolicy='no-referrer')
            # Keep supported toolbar formatting, without CSS URLs or positioning.
            style = str(tag.get('style', ''))
            allowed = []
            for key, value in re.findall(r'([\w-]+)\s*:\s*([^;]+)', style):
                if key == 'text-align' and value.strip() in {'left', 'center', 'right', 'justify'}:
                    allowed.append(f'{key}:{value.strip()}')
                if key in {'color', 'background-color'} and re.fullmatch(r'#[0-9a-fA-F]{3,8}|rgb\([\d ,]+\)', value.strip()):
                    allowed.append(f'{key}:{value.strip()}')
                if key == 'font-size' and re.fullmatch(r'(?:[1-9]|[1-5][0-9]|60)(?:px|pt)', value.strip()):
                    allowed.append(f'{key}:{value.strip()}')
            if allowed:
                attrs['style'] = ';'.join(allowed)
            if tag.name == 'font':
                if str(tag.get('size')) in {'1','2','3','4','5','6','7'}:
                    attrs['size'] = tag['size']
                if tag.get('face') in {'Arial', 'Calibri', 'Georgia'}:
                    attrs['face'] = tag['face']
            tag.attrs = attrs
    return str(soup)


def template_layout():
    path = PROJECT_ROOT / 'template.pptx'
    result = {'available': path.exists(), 'width': 13.333, 'height': 7.5, 'fields': ['title', 'summary', 'insight', 'picture', 'link', 'date', 'team']}
    if path.exists():
        from pptx import Presentation
        deck = Presentation(str(path))
        result.update(width=deck.slide_width / 914400, height=deck.slide_height / 914400)
        result['placeholders'] = []
        layout = next((l for l in deck.slide_layouts if l.name == 'NewsLayout'), None)
        if layout:
            for shape in layout.placeholders:
                result['placeholders'].append({'name': shape.name, 'marker': shape.text if shape.has_text_frame else 'picture', 'x': shape.left / deck.slide_width, 'y': shape.top / deck.slide_height, 'width': shape.width / deck.slide_width, 'height': shape.height / deck.slide_height})
    return result


def export_html(html, title):
    from html import escape
    clean = sanitize_html(html)
    return f'''<!doctype html><html><head><meta charset="utf-8"><title>{escape(title)}</title>
<style>body{{font:16px/1.6 Arial;color:#172438;margin:40px auto;padding:0 32px;max-width:1120px}}h1,h2,h3{{line-height:1.2}}section{{margin-bottom:28px}}section[data-kind=article] img{{max-width:40%;max-height:300px;float:left;margin:0 24px 16px 0}}section[data-kind=article]{{display:flow-root}}img{{max-width:100%}}[data-kind=samsung-impact]{{clear:both;padding:16px 20px;margin:20px 0;background:#f6f9fd;border-left:3px solid #7d9ed0}}[data-kind=page-break]{{break-before:page;clear:both}}@media print{{body{{margin:0;max-width:none}}a{{color:inherit}}}}</style></head><body>{clean}</body></html>'''
