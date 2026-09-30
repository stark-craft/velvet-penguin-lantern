import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from concurrent.futures import ThreadPoolExecutor
from fastapi import FastAPI, HTTPException
from core.storage import JsonStore
from news_scrapper.reports import service
from news_scrapper.reports import analysis
from news_scrapper.reports import impact
from news_scrapper.reports.document import sanitize_html
from news_scrapper.reports.exports import export_snapshot
from news_scrapper.reports.router import router
from tests.asgi_harness import request
from news_scrapper.recommendation.identity import issue_token

class ReportsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.store = JsonStore(Path(self.temp.name)/'reports.json')
        self.patch = patch.object(service, 'STORE', self.store)
        self.patch.start()
    def tearDown(self):
        self.patch.stop();self.temp.cleanup()
    def test_private_revision_drafts_and_history(self):
        first = service.save('a', 'Draft', '<p>Edited text</p>')
        self.assertEqual(service.history('b')['drafts'], [])
        with self.assertRaises(HTTPException) as e:
            service.get_draft('b', first['id'])
        self.assertEqual(e.exception.status_code, 404)
        second = service.save('a', 'Draft', '<p>New edit</p>', first['id'], first['revision'])
        self.assertEqual(second['revision'], 2)
        with self.assertRaises(HTTPException) as e:
            service.save('a','Draft','<p>Lost update</p>',first['id'],1)
        self.assertEqual(e.exception.status_code,409)
        service.record_export('a','Draft','html')
        self.assertFalse(service.history('b')['exports'])
    def test_html_is_sanitized(self):
        text = sanitize_html('<iframe><div>bad</div></iframe><script>alert(1)</script><p onclick="bad()" style="position:fixed;color:#123456">Good<b>bold</b></p><img src="javascript:alert(1)"><a href="javascript:x">Link</a>')
        for evil in ['iframe','script','onclick','position','javascript']:
            self.assertNotIn(evil,text)
        self.assertIn('<b>bold</b>',text)
        self.assertIn('color:#123456',text)
    def configured(self):
        from news_scrapper.adapters import samsung_chat, samsung_web_search
        self.chat=samsung_chat;self.search=samsung_web_search
        self.patches=[patch.object(samsung_chat,'URL','https://example.org/chat'),patch.object(samsung_chat,'CLIENT','test'),patch.object(samsung_chat,'TOKEN','test'),patch.object(samsung_chat,'MODEL_ID','test'),patch.object(samsung_web_search,'ENDPOINT','https://example.org/search'),patch.dict('os.environ',{'SAMSUNG_WEB_SEARCH_TOKEN':'test'}),patch.object(samsung_chat,'call_samsung_chat',return_value={'content':json.dumps({'answer':'Answer','replacement':'Samsung implication'})}),patch.object(samsung_web_search,'call_samsung_web_search_api',return_value={'content':'Evidence','references':[{'link':'https://example.org/article','title':'Source'}]})]
        for p in self.patches:p.start();self.addCleanup(p.stop)
    def test_one_question_two_services_idempotency_and_reset(self):
        self.configured()
        with patch.object(service.time,'time',return_value=100000):
            first=service.ask('a','Question','Context','request-00000001')
            self.assertEqual(first['quota']['remaining'],1)
            self.assertEqual(first['sources'][0]['url'],'https://example.org/article')
            service.ask('a','Question','Context','request-00000001')
            self.assertEqual(self.chat.call_samsung_chat.call_count,1)
            service.ask('a','Another','Context','request-00000002')
            with self.assertRaises(HTTPException) as e:service.ask('a','Third','Context','request-00000003')
            self.assertEqual(e.exception.status_code,429)
            self.assertEqual(e.exception.detail['quota']['reset_at'],121600)
            self.assertEqual(service.quota('b')['remaining'],2)
            replacement=JsonStore(self.store.path)
            with patch.object(service,'STORE',replacement):self.assertEqual(service.quota('a')['remaining'],0)
        with patch.object(service.time,'time',return_value=121600):self.assertEqual(service.quota('a')['remaining'],2)
    def test_atomic_parallel_reservations(self):
        self.configured()
        def attempt(i):
            try:return service.ask('a','Question','Context',f'request-0000000{i}')
            except HTTPException as e:return e.status_code
        with ThreadPoolExecutor(max_workers=8) as pool:results=list(pool.map(attempt,range(8)))
        self.assertEqual(sum(isinstance(x,dict) for x in results),2)
        self.assertEqual(results.count(429),6)
        self.assertEqual(self.chat.call_samsung_chat.call_count,2)
        self.assertEqual(self.search.call_samsung_web_search_api.call_count,2)
    def test_detailed_samsung_impact_uses_chat_and_shared_question_allowance(self):
        self.configured()
        context='Current edited article summary. Existing short Samsung insight.'
        with patch.object(self.search,'ENDPOINT',''):
            result=service.ask('a','Why Samsung?',context,'impact-000000001','samsung-impact')
        self.assertEqual(result['replacement'],'Samsung implication')
        self.assertEqual(result['quota']['remaining'],1)
        self.search.call_samsung_web_search_api.assert_not_called()
        prompt=self.chat.call_samsung_chat.call_args.args[0]
        for expected in ['300–450 words','Opportunities','Risks and uncertainties','Recommended next steps',context,'untrusted evidence']:
            self.assertIn(expected,prompt)
        service.ask('a','Why Samsung?',context,'impact-000000001','samsung-impact')
        self.assertEqual(self.chat.call_samsung_chat.call_count,1)
        service.ask('a','General question',context,'general-00000001')
        with self.assertRaises(HTTPException) as e:
            service.ask('a','Why Samsung?',context,'impact-000000002','samsung-impact')
        self.assertEqual(e.exception.status_code,429)
        self.assertEqual(service.quota('b')['remaining'],2)
    def test_impact_configuration_and_empty_output_never_fake_an_analysis(self):
        self.configured()
        with patch.object(self.chat,'TOKEN',''):
            with self.assertRaises(HTTPException) as e:
                service.ask('a','Why Samsung?','Article','impact-000000001','samsung-impact')
        self.assertEqual(e.exception.status_code,503)
        self.assertEqual(service.quota('a')['remaining'],2)
        self.chat.call_samsung_chat.return_value={'content':json.dumps({'answer':'Explanation','replacement':''})}
        with self.assertRaises(HTTPException) as e:
            service.ask('a','Why Samsung?','Article','impact-000000001','samsung-impact')
        self.assertEqual(e.exception.status_code,502)
        self.assertEqual(service.quota('a')['remaining'],1)
        self.search.call_samsung_web_search_api.assert_not_called()
    def test_impact_route_and_saved_section_contract(self):
        self.configured()
        app=FastAPI();app.include_router(router)
        from news_scrapper.reports import router as module
        headers={'cookie':f'techscout_viewer={issue_token()}'}
        payload={'question':'Why Samsung?','context':'Article','request_id':'impact-000000001','purpose':'samsung-impact'}
        with patch.object(module,'require_capability'):
            result=request(app,'POST','/reports/ask',headers=headers,json_body=payload)
            self.assertEqual(result.status_code,200)
            self.assertEqual(result.json()['replacement'],'Samsung implication')
            payload['purpose']='unsupported'
            self.assertEqual(request(app,'POST','/reports/ask',headers=headers,json_body=payload).status_code,422)
        html='<section data-kind="article"><h2>Article</h2><h3>Summary</h3><p>Facts</p><div data-kind="samsung-impact"><h3>Why this matters to Samsung</h3><div data-kind="impact-body"><p>Detailed strategic relevance</p><p>Recommended next steps</p></div></div></section>'
        saved=service.save('a','Report',html)
        reopened=service.get_draft('a',saved['id'])['html']
        self.assertIn('data-kind="impact-body"',reopened)
        binary,_,_=export_snapshot('html','Report',reopened)
        self.assertIn(b'Detailed strategic relevance',binary)
        binary,_,_=export_snapshot('pdf','Report',reopened)
        from pypdf import PdfReader
        self.assertIn('Recommended next steps',' '.join(p.extract_text() for p in PdfReader(BytesIO(binary)).pages))
    def test_failed_upstream_remains_charged_without_leaking_credentials(self):
        self.configured()
        self.search.call_samsung_web_search_api.side_effect=RuntimeError('secret-token')
        with self.assertRaises(HTTPException) as e:service.ask('a','Question','Context','request-00000001')
        self.assertEqual(e.exception.status_code,502)
        self.assertNotIn('secret-token',str(e.exception.detail))
        self.assertEqual(service.quota('a')['remaining'],1)
    def analysis_articles(self):
        return [{'title':'First development','summary':'Full first summary with evidence.','source':'First publisher','url':'https://example.org/first'},
                {'title':'Second development','summary':'Full second summary with evidence.','source':'Second publisher','url':'https://example.org/second'}]
    def test_automatic_analysis_is_free_even_after_both_ask_questions_are_used(self):
        self.configured()
        service.ask('a','First question','Context','request-00000001')
        service.ask('a','Second question','Context','request-00000002')
        questions=self.store.read()['a']['questions']
        self.chat.call_samsung_chat.return_value={'content':json.dumps({'analysis':'Shared themes: Detailed comparison.\nImplications for Samsung: Inference.'})}
        result=analysis.generate('a',self.analysis_articles())
        self.assertIn('Detailed comparison',result['analysis'])
        self.assertEqual(service.quota('a')['remaining'],0)
        self.assertEqual(self.store.read()['a']['questions'],questions)
        self.assertEqual(self.search.call_samsung_web_search_api.call_count,2)
        self.assertNotIn('quota',result)
    def test_analysis_full_context_fair_bounding_and_no_invented_benchmarks(self):
        articles=self.analysis_articles()
        articles[0]['summary']='A'*19970+'END_OF_FIRST_FULL_SUMMARY'
        articles[1]['summary']='B'*19970+'END_OF_SECOND_FULL_SUMMARY'
        prompt=analysis.analysis_prompt(articles)
        for phrase in ['END_OF_FIRST_FULL_SUMMARY','END_OF_SECOND_FULL_SUMMARY','qualitative comparison','untrusted source data','ALL selected articles']:
            self.assertIn(phrase,prompt)
        many=[{**articles[0],'title':f'Article {i}'} for i in range(20)]
        bounded=analysis.analysis_prompt(many)
        self.assertLess(len(bounded),110000)
        for i in range(20):self.assertIn(f'Article {i}',bounded)
        self.assertIn('"summary_truncated": true',bounded)
        pre_truncated=[{**a,'summary':'Short excerpt','summary_truncated':True} for a in articles]
        self.assertIn('"summary_truncated": true',analysis.analysis_prompt(pre_truncated))
    def test_analysis_cache_is_private_durable_expires_and_tracks_source_changes(self):
        self.configured()
        self.chat.call_samsung_chat.return_value={'content':json.dumps({'analysis':'Comparison'})}
        with patch.object(analysis.time,'time',return_value=100000):
            first=analysis.generate('a',self.analysis_articles())
            self.assertFalse(first['cached'])
            with patch.object(service,'STORE',JsonStore(self.store.path)):
                self.assertTrue(analysis.generate('a',self.analysis_articles())['cached'])
            self.assertEqual(self.chat.call_samsung_chat.call_count,1)
            changed=self.analysis_articles();changed[1]['summary']='Updated second summary'
            analysis.generate('a',changed)
            analysis.generate('b',self.analysis_articles())
            self.assertEqual(self.chat.call_samsung_chat.call_count,3)
            self.assertEqual(service.quota('a')['remaining'],2)
            self.assertNotIn('analyses',service.history('a'))
        with patch.object(analysis.time,'time',return_value=121600):
            analysis.generate('a',self.analysis_articles())
            self.assertEqual(self.chat.call_samsung_chat.call_count,4)
    def test_automatic_impact_is_free_even_after_both_questions_are_used(self):
        self.configured()
        service.ask('a','First question','Context','request-00000001')
        service.ask('a','Second question','Context','request-00000002')
        questions=self.store.read()['a']['questions']
        self.chat.call_samsung_chat.return_value={'content':json.dumps({'answer':'Explanation','replacement':'Strategic relevance: Detailed Samsung opportunities.\nRisks and uncertainties: Evidence gap.'})}
        result=impact.generate('a',self.analysis_articles()[0])
        self.assertIn('Samsung opportunities',result['impact'])
        self.assertEqual(service.quota('a')['remaining'],0)
        self.assertEqual(self.store.read()['a']['questions'],questions)
        self.assertEqual(self.search.call_samsung_web_search_api.call_count,2)
        self.assertNotIn('quota',result)
    def test_automatic_impact_uses_full_evidence_and_discloses_truncation(self):
        self.configured()
        article=self.analysis_articles()[0]
        article['summary']='A'*19965+'END_OF_FULL_SAMSUNG_SUMMARY'
        impact.generate('a',article)
        prompt=self.chat.call_samsung_chat.call_args.args[0]
        for phrase in ['END_OF_FULL_SAMSUNG_SUMMARY','300–450 words','Opportunities','Risks and uncertainties','Recommended next steps','untrusted evidence','First publisher']:
            self.assertIn(phrase,prompt)
        self.search.call_samsung_web_search_api.assert_not_called()
        expanded={**article,'summary':'"'*20000}
        context=impact.article_context(expanded)
        self.assertLessEqual(len(context),30000)
        decoded=json.loads(context)
        self.assertTrue(decoded['summary_truncated'])
        self.assertEqual(decoded['source'],'First publisher')
        self.assertIn('missing or truncated summary',impact.impact_prompt(context))
    def test_automatic_impact_cache_is_private_durable_bounded_and_expires(self):
        self.configured()
        article=self.analysis_articles()[0]
        with patch.object(impact.time,'time',return_value=100000):
            self.assertFalse(impact.generate('a',article)['cached'])
            with patch.object(service,'STORE',JsonStore(self.store.path)):
                self.assertTrue(impact.generate('a',article)['cached'])
            self.assertEqual(self.chat.call_samsung_chat.call_count,1)
            impact.generate('b',article)
            impact.generate('a',{**article,'summary':'Updated article evidence'})
            self.assertEqual(self.chat.call_samsung_chat.call_count,3)
            with patch.object(impact,'CACHE_LIMIT',2):
                impact.generate('a',{**article,'summary':'Third source revision'})
                self.assertEqual(len(self.store.read()['a']['impacts']),2)
            self.assertNotIn('impacts',service.history('a'))
        with patch.object(impact.time,'time',return_value=121600):
            self.assertFalse(impact.generate('b',article)['cached'])
        self.assertEqual(service.quota('a')['remaining'],2)
    def test_parallel_automatic_impact_openings_share_one_chat_call(self):
        import threading
        self.configured()
        started=threading.Event();release=threading.Event()
        def reply(prompt):
            started.set();release.wait(2)
            return {'content':json.dumps({'replacement':'One detailed Samsung impact'})}
        self.chat.call_samsung_chat.side_effect=reply
        with ThreadPoolExecutor(max_workers=8) as pool:
            futures=[pool.submit(impact.generate,'a',self.analysis_articles()[0]) for _ in range(8)]
            self.assertTrue(started.wait(2));release.set()
            results=[f.result() for f in futures]
        self.assertEqual(self.chat.call_samsung_chat.call_count,1)
        self.assertTrue(all(r['impact']=='One detailed Samsung impact' for r in results))
        self.assertEqual(service.quota('a')['remaining'],2)
        self.search.call_samsung_web_search_api.assert_not_called()
    def test_missing_or_failed_automatic_impact_does_not_charge_or_cache(self):
        self.configured()
        with patch.object(self.chat,'TOKEN',''):
            with self.assertRaises(HTTPException) as e:impact.generate('a',self.analysis_articles()[0])
        self.assertEqual(e.exception.status_code,503)
        for content in [json.dumps({'replacement':''}),json.dumps({'replacement':{'invalid':'object'}}),'not JSON']:
            self.chat.call_samsung_chat.return_value={'content':content}
            with self.assertRaises(HTTPException) as e:impact.generate('a',self.analysis_articles()[0])
            self.assertEqual(e.exception.status_code,502)
        self.chat.call_samsung_chat.side_effect=RuntimeError('secret-token')
        with self.assertRaises(HTTPException) as e:impact.generate('a',self.analysis_articles()[0])
        self.assertNotIn('secret-token',str(e.exception.detail))
        self.assertEqual(service.quota('a')['remaining'],2)
        self.assertNotIn('impacts',self.store.read().get('a',{}))
        with self.assertRaises(HTTPException) as e:impact.generate('a',{'title':'Missing summary','summary':''})
        self.assertEqual(e.exception.status_code,422)
    def test_automatic_impact_route_is_private_and_permission_guarded(self):
        self.configured()
        app=FastAPI();app.include_router(router)
        from news_scrapper.reports import router as module
        headers={'cookie':f'techscout_viewer={issue_token()}'}
        with patch.object(module,'require_capability') as permission:
            response=request(app,'POST','/reports/impact',headers=headers,json_body=self.analysis_articles()[0])
            self.assertEqual(response.status_code,200)
            self.assertEqual(response.json()['impact'],'Samsung implication')
            self.assertEqual(response.headers['cache-control'],'no-store')
            self.assertEqual(permission.call_args.args[1],'review.news.submit')
            self.assertEqual(request(app,'POST','/reports/impact',headers=headers,json_body={'title':'A','summary':'x'*20001}).status_code,422)
        with patch.object(module,'require_capability',side_effect=HTTPException(403,'Not allowed')):
            self.assertEqual(request(app,'POST','/reports/impact',headers=headers,json_body=self.analysis_articles()[0]).status_code,403)
    def test_parallel_analysis_openings_share_one_chat_call(self):
        import threading
        self.configured()
        started=threading.Event();release=threading.Event()
        def reply(prompt):
            started.set();release.wait(2)
            return {'content':json.dumps({'analysis':'Single generated comparison'})}
        self.chat.call_samsung_chat.side_effect=reply
        with ThreadPoolExecutor(max_workers=8) as pool:
            futures=[pool.submit(analysis.generate,'a',self.analysis_articles()) for _ in range(8)]
            self.assertTrue(started.wait(2));release.set()
            results=[f.result() for f in futures]
        self.assertEqual(self.chat.call_samsung_chat.call_count,1)
        self.assertTrue(all(r['analysis']=='Single generated comparison' for r in results))
        self.assertEqual(service.quota('a')['remaining'],2)
        self.search.call_samsung_web_search_api.assert_not_called()
    def test_missing_or_failed_automatic_analysis_does_not_charge_or_cache(self):
        self.configured()
        with patch.object(self.chat,'TOKEN',''):
            with self.assertRaises(HTTPException) as e:analysis.generate('a',self.analysis_articles())
        self.assertEqual(e.exception.status_code,503)
        for content in [json.dumps({'analysis':''}),json.dumps({'analysis':{'invalid':'object'}}),'not JSON']:
            self.chat.call_samsung_chat.return_value={'content':content}
            with self.assertRaises(HTTPException) as e:analysis.generate('a',self.analysis_articles())
            self.assertEqual(e.exception.status_code,502)
        self.assertEqual(service.quota('a')['remaining'],2)
        self.assertNotIn('analyses',self.store.read().get('a',{}))
        with self.assertRaises(HTTPException) as e:analysis.generate('a',[{'title':'One','summary':'Single'}])
        self.assertEqual(e.exception.status_code,422)
    def test_automatic_analysis_route_permissions_and_section_exports(self):
        self.configured()
        self.chat.call_samsung_chat.return_value={'content':json.dumps({'analysis':'Compared article evidence'})}
        app=FastAPI();app.include_router(router)
        from news_scrapper.reports import router as module
        headers={'cookie':f'techscout_viewer={issue_token()}'}
        with patch.object(module,'require_capability'):
            response=request(app,'POST','/reports/analysis',headers=headers,json_body={'articles':self.analysis_articles()})
            self.assertEqual(response.status_code,200)
            self.assertEqual(response.json()['analysis'],'Compared article evidence')
            self.assertEqual(response.headers['cache-control'],'no-store')
            self.assertEqual(request(app,'POST','/reports/analysis',headers=headers,json_body={'articles':self.analysis_articles()[:1]}).status_code,422)
        with patch.object(module,'require_capability',side_effect=HTTPException(403,'Forbidden')):
            self.assertEqual(request(app,'POST','/reports/analysis',headers=headers,json_body={'articles':self.analysis_articles()}).status_code,403)
        html='<section data-kind="overview"><h2>Cross-article analysis</h2><div data-kind="cross-analysis"><p>My edited comparison</p></div></section>'
        saved=service.save('a','Report',html)
        self.assertIn('data-kind="cross-analysis"',saved['html'])
        binary,_,_=export_snapshot('docx','Report',saved['html'])
        from docx import Document
        self.assertIn('My edited comparison',' '.join(p.text for p in Document(BytesIO(binary)).paragraphs))
    def test_snapshot_exports_and_pdf_pagination(self):
        html='<section data-kind="overview"><h1>Report</h1><p>My edited wording <b>bold</b></p><div>Typing with Enter</div><p>=DANGEROUS()</p><hr data-kind="page-break"><p>Second page</p></section>'
        binary,_,_=export_snapshot('html','Title',html)
        self.assertIn(b'My edited wording',binary)
        binary,_,_=export_snapshot('pdf','Title',html)
        from pypdf import PdfReader
        pdf=PdfReader(BytesIO(binary));self.assertEqual(len(pdf.pages),2)
        self.assertIn('My edited wording',pdf.pages[0].extract_text())
        self.assertIn('Typing with Enter',pdf.pages[0].extract_text())
        binary,_,_=export_snapshot('docx','Title',html)
        from docx import Document
        doc=Document(BytesIO(binary));self.assertIn('My edited wording',' '.join(p.text for p in doc.paragraphs))
        binary,_,_=export_snapshot('xlsx','Title',html)
        from openpyxl import load_workbook
        book=load_workbook(BytesIO(binary));cells=list(book.active.values)
        self.assertTrue(any(row[1]=="'=DANGEROUS()" for row in cells))
    def test_local_replacement_image_survives_save_html_pdf_and_word(self):
        import base64
        from PIL import Image
        from pypdf import PdfReader
        from docx import Document
        picture=BytesIO()
        Image.new('RGB',(20,20),'navy').save(picture,format='PNG')
        url='data:image/png;base64,'+base64.b64encode(picture.getvalue()).decode()
        saved=service.save('owner','Local picture',f'<section data-kind="article"><h2>News</h2><img src="{url}" data-report-selected="true" role="button" tabindex="0"><p>Edited summary</p></section>')
        html=service.get_draft('owner',saved['id'])['html']
        self.assertIn(url,html)
        self.assertNotIn('data-report-selected',html)
        binary,_,_=export_snapshot('html','Report',html)
        self.assertIn(url.encode(),binary)
        binary,_,_=export_snapshot('pdf','Report',html)
        self.assertEqual(len(PdfReader(BytesIO(binary)).pages[0].images),1)
        binary,_,_=export_snapshot('docx','Report',html)
        self.assertEqual(len(Document(BytesIO(binary)).inline_shapes),1)
    def test_ppt_uses_template_markers_and_edited_fields(self):
        from pptx import Presentation
        from news_scrapper.reports import exports
        deck=Presentation()
        layout=deck.slide_layouts[3]
        layout.name='NewsLayout'
        layout.placeholders[0].text='#TITLE'
        layout.placeholders[1].text='#SUMMARY'
        layout.placeholders[2].text='#INSIGHT'
        deck.save(str(Path(self.temp.name)/'template.pptx'))
        html='<section data-kind="overview"><h1>Report</h1><h2>Executive summary</h2><p>Overview edit</p><h2>Cross-article analysis</h2><p>Analysis edit</p></section><section data-kind="article"><h2>Edited headline</h2><p>Publisher date</p><h3>Summary</h3><p>Edited summary</p><div data-kind="samsung-impact"><h3>Why this matters to Samsung</h3><div data-kind="impact-body"><p>Edited insight</p></div></div></section>'
        with patch.object(exports,'PROJECT_ROOT',Path(self.temp.name)):
            data,_,_=export_snapshot('pptx','Report',html)
        result=Presentation(BytesIO(data))
        self.assertEqual(len(result.slides),2)
        text=' '.join(shape.text for slide in result.slides for shape in slide.shapes if shape.has_text_frame)
        for wording in ['Overview edit','Analysis edit','Edited headline','Edited summary','Edited insight']:
            self.assertIn(wording,text)
    def test_missing_ppt_template_is_explicit(self):
        from news_scrapper.reports import exports
        with patch.object(exports,'PROJECT_ROOT',Path(self.temp.name)):
            with self.assertRaises(HTTPException) as e:
                export_snapshot('pptx','Title','<p>Edited</p>')
        self.assertEqual(e.exception.status_code,409)

    def test_route_ownership_and_permissions(self):
        app=FastAPI();app.include_router(router)
        from news_scrapper.reports import router as module
        a={'cookie':f'techscout_viewer={issue_token()}'}
        b={'cookie':f'techscout_viewer={issue_token()}'}
        with patch.object(module,'require_capability'):
            created=request(app,'POST','/reports/drafts',headers=a,json_body={'title':'Mine','html':'<p>Private</p>'})
            self.assertEqual(created.status_code,200)
            denied=request(app,'GET',f"/reports/drafts/{created.json()['id']}",headers=b)
            self.assertEqual(denied.status_code,404)
            self.assertEqual(request(app,'GET','/reports/history',headers=b).json()['drafts'],[])
        with patch.object(module,'require_capability',side_effect=HTTPException(403,'Forbidden')):
            self.assertEqual(request(app,'GET','/reports/status',headers=a).status_code,403)

if __name__=='__main__':unittest.main()
