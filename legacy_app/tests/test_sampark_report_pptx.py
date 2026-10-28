"""Regression tests for the Sampark Create Report -> PowerPoint export."""
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from bs4 import BeautifulSoup
from fastapi import HTTPException
from pptx import Presentation
from pptx.util import Inches

from news_scrapper.reports import exports


ARTICLE_HTML = """
<section data-kind="overview"><h1>Executive summary</h1><p>Not a PPT slide.</p></section>
<section data-kind="article">
  <h2>OpenAI launches a new AI product</h2>
  <p>TechCrunch · 2026-10-28</p>
  <div><h3>Summary</h3>
    <div data-kind="article-summary"><p>First article summary.</p></div>
    <div data-kind="samsung-impact"><h3>Why this matters to Samsung</h3>
      <div data-kind="impact-body"><p>First insight for Samsung.</p></div>
    </div>
    <h3>Relevant teams</h3><p>AI Research</p>
    <p><a href="https://example.org/first">Open original article</a></p>
  </div>
</section>
<section data-kind="article">
  <h2>Second report headline</h2>
  <p>Reuters · 2026-10-27</p>
  <div><h3>Summary</h3>
    <div data-kind="article-summary"><p>Second article summary.</p></div>
    <div data-kind="samsung-impact"><h3>Why this matters to Samsung</h3>
      <div data-kind="impact-body"><p>Second insight.</p></div>
    </div>
    <h3>Relevant teams</h3><p>Device Solutions</p>
    <p><a href="https://example.org/second">Open original article</a></p>
  </div>
</section>
"""


class SamparkPowerPointExportTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.root_patch = patch.object(exports, 'PROJECT_ROOT', self.root)
        self.image_patch = patch.object(exports, 'image_data', return_value=None)
        self.root_patch.start()
        self.image_patch.start()
        self.addCleanup(self.root_patch.stop)
        self.addCleanup(self.image_patch.stop)

    def create_template(self, paired=True):
        """Built-in layouts stand in for article layouts; no proprietary template needed."""
        deck = Presentation()
        first, regular = deck.slide_layouts[1], deck.slide_layouts[2]
        first.name = 'NewsLayoutFirst' if paired else 'UnusedLayout'
        regular.name = 'NewsLayout'
        for layout in (first, regular):
            layout.placeholders[0].text = '#TITLE'
            layout.placeholders[1].text = '#SUMMARY'
        # Verify export discards any design-sample slides from the input template.
        deck.slides.add_slide(first)
        deck.save(str(self.root / 'template.pptx'))

    def test_first_article_uses_first_layout_and_skips_overview(self):
        self.create_template()
        result, mime, extension = exports.export_snapshot('pptx', 'Test report', ARTICLE_HTML)
        self.assertEqual(extension, 'pptx')
        self.assertEqual(mime, 'application/vnd.openxmlformats-officedocument.presentationml.presentation')
        deck = Presentation(BytesIO(result))
        self.assertEqual(len(deck.slides), 2)
        self.assertEqual([s.slide_layout.name for s in deck.slides],
                         ['NewsLayoutFirst', 'NewsLayout'])
        self.assertIn('OpenAI launches', deck.slides[0].placeholders[0].text)
        self.assertIn('Second report', deck.slides[1].placeholders[0].text)
        self.assertNotIn('Executive summary', deck.slides[0].placeholders[0].text)
        self.assertEqual(deck.slides[0].placeholders[1].text, 'First article summary.')
        self.assertEqual(deck.slides[1].placeholders[1].text, 'Second article summary.')
        self.assertEqual(len(Presentation(str(self.root / 'template.pptx')).slides), 1)

    def test_headline_is_linked_and_black_byline_is_not(self):
        self.create_template()
        result, _, _ = exports.export_snapshot('pptx', 'Test report', ARTICLE_HTML)
        deck = Presentation(BytesIO(result))
        paragraph = deck.slides[0].placeholders[0].text_frame.paragraphs[0]
        self.assertEqual(len(paragraph.runs), 2)
        headline, byline = paragraph.runs
        self.assertEqual(headline.text, 'OpenAI launches a new AI product')
        self.assertEqual(headline.hyperlink.address, 'https://example.org/first')
        self.assertEqual(headline.font.name, 'Calibri')
        self.assertTrue(headline.font.bold)
        self.assertEqual(str(headline.font.color.rgb), '0070C0')
        self.assertEqual(byline.text, ' By TechCrunch')
        self.assertIsNone(byline.hyperlink.address)
        self.assertFalse(byline.font.bold)
        self.assertEqual(str(byline.font.color.rgb), '000000')

    def test_summary_and_long_title_are_fitted_without_widening_boxes(self):
        self.create_template()
        deck = Presentation(str(self.root / 'template.pptx'))
        for layout in (deck.slide_layouts[1], deck.slide_layouts[2]):
            layout.placeholders[0].width = Inches(9.519)
            layout.placeholders[0].height = Inches(.703)
            layout.placeholders[1].width = Inches(7.481)
            layout.placeholders[1].height = Inches(2.148)
        deck.save(str(self.root / 'template.pptx'))
        long_title = 'A significantly extended story with multiple technologies and important considerations ' * 7
        long_summary = 'A significant technical development was reported. ' * 120
        html = ARTICLE_HTML.replace('OpenAI launches a new AI product', long_title).replace(
            'First article summary.', long_summary
        )
        result, _, _ = exports.export_snapshot('pptx', 'Test report', html)
        slide = Presentation(BytesIO(result)).slides[0]
        title = slide.placeholders[0]
        summary = slide.placeholders[1]
        self.assertIn('By TechCrunch', title.text)
        self.assertIn('…', title.text)
        self.assertTrue(exports._pptx_fits(title.text, title, title.text_frame.paragraphs[0].runs[0].font.size.pt, True))
        self.assertTrue(summary.text.endswith('…'))
        self.assertTrue(exports._pptx_fits(summary.text, summary, 16))
        self.assertAlmostEqual(summary.width / 914400, 7.481, places=2)

    def test_extract_article_uses_edited_snapshot_not_overview_or_link_as_team(self):
        article = BeautifulSoup(ARTICLE_HTML, 'html.parser').select_one('section[data-kind="article"]')
        parsed = exports._pptx_article(article)
        self.assertEqual(parsed['source'], 'TechCrunch')
        self.assertEqual(parsed['date'], '2026-10-28')
        self.assertEqual(parsed['insight'], 'First insight for Samsung.')
        self.assertEqual(parsed['team'], 'AI Research')
        self.assertEqual(parsed['link'], 'https://example.org/first')

    def test_missing_first_layout_is_a_clear_configuration_error(self):
        self.create_template(paired=False)
        with self.assertRaises(HTTPException) as captured:
            exports.export_snapshot('pptx', 'Test report', ARTICLE_HTML)
        self.assertEqual(captured.exception.status_code, 409)
        self.assertIn('NewsLayoutFirst', captured.exception.detail)

    def test_zero_articles_are_not_exported_as_overview_slides(self):
        self.create_template()
        with self.assertRaises(HTTPException) as captured:
            exports.export_snapshot(
                'pptx', 'Test report',
                '<section data-kind="overview"><h1>Executive summary</h1></section>',
            )
        self.assertEqual(captured.exception.status_code, 422)


if __name__ == '__main__':
    unittest.main()
