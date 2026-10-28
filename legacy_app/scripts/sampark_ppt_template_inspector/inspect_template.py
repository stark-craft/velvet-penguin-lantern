#!/usr/bin/env python3
"""Read-only Sampark PPTX template inspector.

Never copies embedded images, slide text, speaker notes, or the PowerPoint file
into its outputs. Generates an inspection JSON, a schematic HTML map and a ZIP.
The map is NOT a pixel-accurate PowerPoint rendering.

Requires only python-pptx (already listed in the TechScout backend requirements).
Run from an isolated local folder; do not put proprietary PPTX files in Git.
"""

import argparse
from collections import Counter
from datetime import datetime, timezone
from html import escape
import json
from pathlib import Path
import sys
import zipfile
import xml.etree.ElementTree as ET

EMU_PER_INCH = 914400
EMU_PER_PT = 12700
TAGS = {
    "title": "#TITLE",
    "summary": "#SUMMARY",
    "insight": "#INSIGHT",
    "link": "#LINK",
    "date": "#DATE_HERE",
    "team": "#Targated_SRID_Team",
}
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def round_or_none(value, digits=3):
    try:
        return round(float(value), digits) if value is not None else None
    except (TypeError, ValueError, OverflowError):
        return None


def inches(value):
    return round_or_none(float(value) / EMU_PER_INCH if value is not None else None, 4)


def points(value):
    return round_or_none(float(value) / EMU_PER_PT if value is not None else None, 2)


def safe_get(call, default=None):
    try:
        return call()
    except (AttributeError, TypeError, ValueError, KeyError, IndexError, OSError):
        return default


def enum_name(value):
    if value is None:
        return None
    text = str(value)
    return text.split(" (", 1)[0]


def color_value(value):
    """Report explicit RGB/theme colors only; no private image data."""
    if value is None:
        return None
    kind = safe_get(lambda: enum_name(value.type))
    if kind is None:
        return None
    out = {"kind": kind}
    rgb = safe_get(lambda: str(value.rgb))
    theme = safe_get(lambda: enum_name(value.theme_color))
    brightness = safe_get(lambda: value.brightness)
    if rgb:
        out["rgb"] = rgb
    if theme:
        out["theme_color"] = theme
    if brightness is not None:
        out["brightness"] = round_or_none(brightness, 4)
    return out


def font_info(font):
    if font is None:
        return {}
    size = safe_get(lambda: font.size)
    result = {
        "name_explicit": safe_get(lambda: font.name),
        "size_pt_explicit": points(size),
        "bold_explicit": safe_get(lambda: font.bold),
        "italic_explicit": safe_get(lambda: font.italic),
        "underline_explicit": safe_get(lambda: font.underline),
        "color_explicit": safe_get(lambda: color_value(font.color)),
    }
    return {key: value for key, value in result.items() if value is not None}


def dimensions(shape, slide_w, slide_h):
    x = safe_get(lambda: int(shape.left), 0)
    y = safe_get(lambda: int(shape.top), 0)
    w = safe_get(lambda: int(shape.width), 0)
    h = safe_get(lambda: int(shape.height), 0)
    return {
        "x_in": inches(x), "y_in": inches(y),
        "width_in": inches(w), "height_in": inches(h),
        "x_pct": round_or_none(x * 100 / slide_w),
        "y_pct": round_or_none(y * 100 / slide_h),
        "width_pct": round_or_none(w * 100 / slide_w),
        "height_pct": round_or_none(h * 100 / slide_h),
    }


def matches_marker(text):
    # The actual slide text is used in memory ONLY to match these public tokens.
    return [token for token in TAGS.values() if token.lower() in text.lower()]


def paragraph_info(paragraph):
    text = safe_get(lambda: paragraph.text, "") or ""
    line = safe_get(lambda: paragraph.line_spacing)
    if isinstance(line, (int, float)):
        line_value = points(line) if abs(line) > 10 else round_or_none(line)
    else:
        line_value = None
    runs = list(safe_get(lambda: paragraph.runs, []) or [])
    styles = []
    for run in runs:
        style = font_info(run.font)
        if style and style not in styles and len(styles) < 15:
            styles.append(style)
    return {
        "characters": len(text),
        "line_breaks": text.count("\n"),
        "level": safe_get(lambda: paragraph.level),
        "alignment": safe_get(lambda: enum_name(paragraph.alignment)),
        "line_spacing_explicit": line_value,
        "space_before_pt_explicit": points(safe_get(lambda: paragraph.space_before)),
        "space_after_pt_explicit": points(safe_get(lambda: paragraph.space_after)),
        "style_samples_without_text": styles,
        "run_count": len(runs),
    }


def shape_info(shape, ordinal, slide_w, slide_h, depth=0):
    record = {
        "object": "shape_" + str(ordinal),
        "shape_type": safe_get(lambda: enum_name(shape.shape_type)),
        "shape_id": safe_get(lambda: shape.shape_id),
        "geometry": dimensions(shape, slide_w, slide_h),
        "rotation_degrees": round_or_none(safe_get(lambda: shape.rotation), 2),
    }
    if safe_get(lambda: shape.is_placeholder, False):
        fmt = safe_get(lambda: shape.placeholder_format)
        record["placeholder"] = {
            "index": safe_get(lambda: fmt.idx),
            "type": safe_get(lambda: enum_name(fmt.type)),
        }
    fill_type = safe_get(lambda: enum_name(shape.fill.type))
    if fill_type:
        record["fill"] = {"type": fill_type}
        if "SOLID" in fill_type:
            record["fill"]["color"] = safe_get(lambda: color_value(shape.fill.fore_color))
    line = safe_get(lambda: shape.line)
    if line:
        record["outline"] = {
            "width_pt": points(safe_get(lambda: line.width)),
            "color": safe_get(lambda: color_value(line.color)),
        }
    if safe_get(lambda: shape.has_text_frame, False):
        tf = shape.text_frame
        txt = safe_get(lambda: shape.text, "") or ""
        record["recognized_tokens_only"] = matches_marker(txt)
        record["text_statistics_no_content"] = {
            "characters": len(txt),
            "paragraphs": len(tf.paragraphs),
            "line_breaks": txt.count("\n"),
        }
        inner_w = max(0, safe_get(lambda: shape.width, 0) -
                       (safe_get(lambda: tf.margin_left, 0) or 0) -
                       (safe_get(lambda: tf.margin_right, 0) or 0))
        inner_h = max(0, safe_get(lambda: shape.height, 0) -
                       (safe_get(lambda: tf.margin_top, 0) or 0) -
                       (safe_get(lambda: tf.margin_bottom, 0) or 0))
        record["text_frame"] = {
            "left_margin_in": inches(safe_get(lambda: tf.margin_left)),
            "right_margin_in": inches(safe_get(lambda: tf.margin_right)),
            "top_margin_in": inches(safe_get(lambda: tf.margin_top)),
            "bottom_margin_in": inches(safe_get(lambda: tf.margin_bottom)),
            "usable_width_in": inches(inner_w),
            "usable_height_in": inches(inner_h),
            "wrap_explicit": safe_get(lambda: tf.word_wrap),
            "autofit_explicit": safe_get(lambda: enum_name(tf.auto_size)),
            "vertical_alignment_explicit": safe_get(lambda: enum_name(tf.vertical_anchor)),
        }
        record["paragraph_styles"] = [
            paragraph_info(p) for p in list(tf.paragraphs)[:100]
        ]
        if len(tf.paragraphs) > 100:
            record["paragraph_styles_truncated"] = True
    if record.get("shape_type", "").startswith("PICTURE"):
        img = safe_get(lambda: shape.image)
        if img is not None:
            width_height = safe_get(lambda: img.size)
            record["image_metadata_only"] = {
                "extension": safe_get(lambda: img.ext),
                "natural_width_px": width_height[0] if width_height else None,
                "natural_height_px": width_height[1] if width_height else None,
                "pixels_not_exported": True,
            }
    if safe_get(lambda: shape.has_table, False):
        table = shape.table
        record["table_dimensions_no_cell_text"] = {
            "rows": len(table.rows),
            "columns": len(table.columns),
        }
    if safe_get(lambda: shape.has_chart, False):
        record["chart_data_exported"] = False
    # Group children have a different coordinate space; preserve their raw
    # dimensions and flag them as such instead of drawing at incorrect positions.
    children = safe_get(lambda: shape.shapes)
    if children is not None and depth < 3:
        record["children_in_group_local_coordinates"] = [
            shape_info(child, i + 1, slide_w, slide_h, depth + 1)
            for i, child in enumerate(children)
        ]
    return record


def collect_shapes(shapes, width, height):
    return [shape_info(shape, i + 1, width, height)
            for i, shape in enumerate(shapes)]


def theme_info(filename):
    """Theme font/color maps from OOXML; these can be inherited by shapes."""
    themes = []
    with zipfile.ZipFile(filename) as archive:
        names = sorted(name for name in archive.namelist()
                       if name.startswith("ppt/theme/") and name.endswith(".xml"))
        for name in names[:40]:
            root = ET.fromstring(archive.read(name))
            font_scheme = root.find(".//a:fontScheme", NS)
            clr_scheme = root.find(".//a:clrScheme", NS)
            colors = {}
            fonts = {}
            if font_scheme is not None:
                for key in ("majorFont", "minorFont"):
                    latin = font_scheme.find("a:" + key + "/a:latin", NS)
                    fonts[key] = latin.attrib.get("typeface") if latin is not None else None
            if clr_scheme is not None:
                for entry in clr_scheme:
                    child = next(iter(entry), None)
                    if child is not None:
                        value = child.get("val") or child.get("lastClr")
                        if value:
                            colors[entry.tag.rsplit("}", 1)[-1]] = value
            themes.append({"theme_number": len(themes) + 1, "fonts": fonts, "colors": colors})
    return themes


def slide_layout_info(layout, index, sw, sh):
    return {
        "index": index,
        "layout_name": layout.name,
        "shapes": collect_shapes(layout.shapes, sw, sh),
    }


def inspection_report(pptx, max_slides):
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError(
            "python-pptx was not found. Use your project's Python, e.g. "
            r"C:\scrappyV3\venv\Scripts\python.exe, or install python-pptx."
        ) from exc
    presentation = Presentation(str(pptx))
    sw, sh = int(presentation.slide_width), int(presentation.slide_height)
    masters = []
    for master_index, master in enumerate(presentation.slide_masters, 1):
        masters.append({
            "index": master_index,
            "shapes": collect_shapes(master.shapes, sw, sh),
            "layouts": [
                slide_layout_info(layout, idx, sw, sh)
                for idx, layout in enumerate(master.slide_layouts)
            ],
        })
    slides = []
    for index, slide in enumerate(presentation.slides):
        if index >= max_slides:
            break
        slides.append({
            "number": index + 1,
            "layout_name": slide.slide_layout.name,
            "shapes": collect_shapes(slide.shapes, sw, sh),
        })
    layouts = [layout for master in masters for layout in master["layouts"]]
    matching = [layout for layout in layouts if layout["layout_name"] == "NewsLayout"]
    required = list(TAGS.values())
    found = set()
    for layout in matching:
        for shape in layout["shapes"]:
            found.update(shape.get("recognized_tokens_only", []))
    has_picture = any(
        shape.get("placeholder", {}).get("type", "").startswith("PICTURE")
        for layout in matching for shape in layout["shapes"]
    )
    return {
        "report_type": "sampark_pptx_structural_inspection",
        "privacy": {
            "slide_text_exported": False,
            "speaker_notes_exported": False,
            "image_pixels_exported": False,
            "original_powerpoint_exported": False,
            "original_file_name_exported": False,
            "shape_names_exported": False,
            "caution": "Dimensions, theme colors, font names and layout names may still be sensitive. Review before sharing.",
        },
        "generated_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "slide_size": {
            "width_in": inches(sw),
            "height_in": inches(sh),
            "width_emu": sw,
            "height_emu": sh,
            "aspect_ratio": round_or_none(sw / sh, 5),
        },
        "total_slide_count": len(presentation.slides),
        "inspected_slide_count": len(slides),
        "master_count": len(masters),
        "layout_count": len(layouts),
        "theme": theme_info(pptx),
        "sampark_legacy_template_contract": {
            "expected_layout_name": "NewsLayout",
            "layout_found": bool(matching),
            "known_marker_tokens": required,
            "markers_found_in_layout": sorted(found),
            "markers_not_found_in_layout": [t for t in required if t not in found],
            "picture_placeholder_found": has_picture,
            "note": "A new template may intentionally use different names or ordinary textboxes. This is a diagnostic, not a failure.",
        },
        "masters": masters,
        "sample_slides": slides,
        "notes": [
            "Geometry is from PowerPoint OOXML in inches and slide-percent.",
            "Fonts listed on runs are explicit overrides, not guaranteed effective/inherited values.",
            "Inherited formatting may come from slide layout, master, theme or PowerPoint defaults.",
            "A diagram does not reproduce embedded artwork, fonts, smart art, image crop or theme effects.",
            "Textbox capacity needs later validation with PowerPoint's renderer. Word counts alone do not guarantee fit.",
        ],
    }


def draw_shape(shape):
    g = shape["geometry"]
    left = round_or_none(g.get("x_pct"), 2) or 0
    top = round_or_none(g.get("y_pct"), 2) or 0
    width = max(0, round_or_none(g.get("width_pct"), 2) or 0)
    height = max(0, round_or_none(g.get("height_pct"), 2) or 0)
    kind = shape.get("shape_type", "") or ""
    tokens = shape.get("recognized_tokens_only", [])
    ph = shape.get("placeholder")
    if tokens:
        label = ", ".join(tokens)
    elif ph:
        label = "placeholder " + str(ph.get("index"))
    elif kind.startswith("PICTURE"):
        label = "Picture"
    elif "text_frame" in shape:
        label = "Text area"
    else:
        label = kind.title()[:24]
    category = "picture" if "PICTURE" in kind else (
        "text" if "text_frame" in shape else "other"
    )
    dims = (
        str(g.get("width_in")) + " x " + str(g.get("height_in")) + " in"
    )
    title = escape(label + " | " + dims, quote=True)
    label = escape(label)
    return (
        '<div class="shape ' + category + '" style="left:' + str(left) +
        '%;top:' + str(top) + '%;width:' + str(width) +
        '%;height:' + str(height) + '%" title="' + title +
        '"><span>' + label + '</span></div>'
    )


def make_page(title, shapes, aspect):
    shapes_html = "\n".join(draw_shape(shape) for shape in shapes)
    return (
        '<section class="card"><h2>' + escape(title) +
        '</h2><div class="slide" style="aspect-ratio:' + str(aspect) +
        '"><div class="gridlines"></div>' + shapes_html +
        '</div><p class="hint">Approximate shape geometry only. '
        'Text and image content have been intentionally removed.</p></section>'
    )


def html_preview(report):
    aspect = report["slide_size"]["aspect_ratio"] or 16 / 9
    panels = []
    for master in report["masters"]:
        panels.append(make_page(
            "Master " + str(master["index"]), master["shapes"], aspect
        ))
        for layout in master["layouts"]:
            panels.append(make_page(
                "Layout: " + layout["layout_name"], layout["shapes"], aspect
            ))
    for slide in report["sample_slides"]:
        panels.append(make_page(
            "Sample slide " + str(slide["number"]) + " (" +
            slide["layout_name"] + ")", slide["shapes"], aspect
        ))
    contract = report["sampark_legacy_template_contract"]
    contract_str = ("NewsLayout found: " + str(contract["layout_found"]) +
                    " | Missing old marker tokens: " +
                    ", ".join(contract["markers_not_found_in_layout"]))
    return """<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Sampark template structural map (redacted)</title>
<style>
:root{font-family:system-ui,Arial,sans-serif;color:#142334;background:#eef2f7}
body{margin:0;padding:25px;max-width:1200px;margin:auto}
h1{font-size:26px;margin-bottom:8px}h2{font-size:16px}
p{line-height:1.5}.warning{border:1px solid #c8d6e7;padding:12px;border-radius:9px;background:white}
.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(320px,1fr));gap:20px}
.card{background:white;padding:16px;border:1px solid #d6deea;border-radius:12px;min-width:0}
.slide{position:relative;width:100%;background:#fff;border:1px solid #b1becb;overflow:hidden}
.gridlines{position:absolute;inset:0;background-image:linear-gradient(90deg,transparent 49.8%,#f0f4f8 50%,transparent 50.2%),
linear-gradient(transparent 49.8%,#f0f4f8 50%,transparent 50.2%);pointer-events:none}
.shape{position:absolute;overflow:hidden;box-sizing:border-box;border:1px solid #7b8797;background:#d8e1eb55}
.shape.text{background:#c2d9fa55;border-color:#3873b5}
.shape.picture{background:#d1f1df88;border-color:#319267}
.shape span{display:block;font-size:10px;line-height:1.3;word-wrap:break-word;padding:2px}
.hint{font-size:11px;color:#637184}.cards{margin-top:20px}
</style></head><body>
<h1>Sampark PowerPoint template — structural map</h1>
<p class="warning"><b>Content-redacted:</b> no slide text, notes, or image pixels.
This is NOT a PowerPoint render. Review the JSON and HTML for company-sensitive
design details before sharing. Do not upload the original PPTX.</p>
<p>""" + escape(contract_str) + """</p><div class="cards">""" + (
        "\n".join(panels)
    ) + """</div></body></html>"""


def create_outputs(report, output_dir):
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "template_structure.json"
    html_path = output_dir / "template_geometry.html"
    zip_path = output_dir / "sampark_template_inspection_REVIEW_BEFORE_SHARING.zip"
    json_path.write_text(
        json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    html_path.write_text(html_preview(report), encoding="utf-8")
    readme = output_dir / "REVIEW_BEFORE_SHARING.txt"
    readme.write_text(
        "This inspection does NOT include the original PPTX, slide text, notes "
        "or embedded image pixels.\n"
        "However layout names, dimensions, fonts, color values and geometry may "
        "still be confidential. Review them under Samsung policies before sharing.\n"
        "The HTML is a schematic, not a pixel-perfect preview.\n",
        encoding="utf-8",
    )
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for file in (json_path, html_path, readme):
            archive.write(file, arcname=file.name)
    return json_path, html_path, zip_path


def main():
    parser = argparse.ArgumentParser(
        description="Inspect a local PPTX with content-redacted JSON and HTML."
    )
    parser.add_argument("template", type=Path, help="Path to new .pptx template")
    parser.add_argument(
        "--out", type=Path, required=True, help="Folder for JSON/HTML/ZIP outputs"
    )
    parser.add_argument(
        "--max-slides", type=int, default=8,
        help="Only inspect up to this many sample slides (default: 8)."
    )
    args = parser.parse_args()
    if not args.template.is_file() or args.template.suffix.lower() != ".pptx":
        parser.error("Template must exist and have .pptx extension.")
    if args.max_slides < 0 or args.max_slides > 100:
        parser.error("--max-slides must be between 0 and 100.")
    try:
        if args.template.resolve().parent == args.out.resolve():
            raise ValueError("Use a separate output folder from the original PPTX.")
        report = inspection_report(args.template, args.max_slides)
        paths = create_outputs(report, args.out)
    except (RuntimeError, ValueError, OSError, zipfile.BadZipFile, ET.ParseError) as exc:
        print("Inspection failed: " + str(exc), file=sys.stderr)
        return 2
    print("\nInspection finished. The ORIGINAL PPTX was not modified.")
    for path in paths:
        print("  " + str(path))
    print("\nReview the ZIP for Samsung-sensitive design information before sharing.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
